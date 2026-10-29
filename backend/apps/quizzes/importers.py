"""Savollarni Excel shablondan import qilish (FT-6.3).

Shablon ustunlari (1-qator sarlavha): Mavzu | Daraja | Savol | A | B | C | D | To'g'ri | Izoh
  Mavzu   — mavzu kaliti yoki nomi («Kepler», «Linza» ...)
  Daraja  — Asosiy / Chuqur
  To'g'ri — A, B, C yoki D
"""
from openpyxl import load_workbook

from .models import Level, Question, Status, Topic


def import_questions_xlsx(f, user) -> tuple[int, list[str]]:
    if not (f.name or "").lower().endswith(".xlsx"):
        return 0, ["Faqat .xlsx fayl qabul qilinadi"]
    ws = load_workbook(f, read_only=True, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    topics = {t.key.lower(): t for t in Topic.objects.all()}
    topics.update({t.title_uz.lower(): t for t in Topic.objects.all()})
    created, errors = 0, []
    status = Status.APPROVED if user.is_platform_admin else Status.DRAFT
    for n, row in enumerate(rows[1:], start=2):
        cells = [("" if c is None else str(c).strip()) for c in (list(row) + [None] * 9)[:9]]
        if not any(cells):
            continue
        topic_name, level, text, a, b, c, d, correct, expl = cells
        topic = topics.get(topic_name.lower())
        if not topic:
            errors.append(f"{n}-qator: mavzu «{topic_name}» topilmadi")
            continue
        lvl = Level.CHUQUR if level.lower().startswith("ch") else Level.ASOSIY
        opts = [a, b, c, d]
        if not text or not all(opts):
            errors.append(f"{n}-qator: savol yoki variantlar to'liq emas")
            continue
        if len({o.lower() for o in opts}) != 4:
            errors.append(f"{n}-qator: variantlar takrorlanmoqda")
            continue
        letter = correct.upper()[:1]
        if letter not in "ABCD" or not letter:
            errors.append(f"{n}-qator: to'g'ri javob A, B, C yoki D bo'lishi kerak")
            continue
        Question.objects.create(topic=topic, level=lvl, text_uz=text, options=[{"uz": o, "ru": ""} for o in opts],
                                correct_index="ABCD".index(letter), explanation_uz=expl, author=user,
                                status=status)
        created += 1
    return created, errors
