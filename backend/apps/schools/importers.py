"""O'quvchilar ro'yxatini CSV yoki Excel fayldan o'qish (FT-1.2)."""
import csv
import io

from rest_framework.exceptions import ValidationError

MAX_ROWS = 60
_LAST = {"familiya", "фамилия", "last_name", "lastname"}
_FIRST = {"ism", "имя", "first_name", "firstname", "ismi"}


def _rows_from_table(rows):
    rows = [[(c or "").strip() if isinstance(c, str) else ("" if c is None else str(c)) for c in r]
            for r in rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        return []
    head = [h.lower() for h in rows[0]]
    li = next((i for i, h in enumerate(head) if h in _LAST), None)
    fi = next((i for i, h in enumerate(head) if h in _FIRST), None)
    if li is not None and fi is not None:
        body = rows[1:]
    else:
        li, fi, body = 0, 1, rows  # sarlavhasiz: 1-ustun familiya, 2-ustun ism
    out = []
    for r in body:
        last = r[li] if li < len(r) else ""
        first = r[fi] if fi < len(r) else ""
        if not first and " " in last:  # "Aliyev Vali" bitta katakda
            last, first = last.split(" ", 1)
        if last or first:
            out.append({"last_name": last[:60], "first_name": first[:60]})
    return out


def parse_students_file(f) -> list[dict]:
    name = (f.name or "").lower()
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        wb = load_workbook(f, read_only=True, data_only=True)
        rows = list(wb.active.iter_rows(values_only=True))
    elif name.endswith(".csv") or name.endswith(".txt"):
        raw = f.read()
        text = raw.decode("utf-8-sig", errors="replace")
        dialect = csv.Sniffer().sniff(text[:2000], delimiters=";,\t") if text.strip() else csv.excel
        rows = list(csv.reader(io.StringIO(text), dialect))
    else:
        raise ValidationError({"file": ["Faqat .xlsx yoki .csv fayl qabul qilinadi"]})
    students = _rows_from_table(rows)
    if not students:
        raise ValidationError({"file": ["Faylda o'quvchilar topilmadi"]})
    return students


def check_limit(students):
    if len(students) > MAX_ROWS:
        raise ValidationError(f"Bir martada ko'pi bilan {MAX_ROWS} ta o'quvchi qo'shish mumkin")
