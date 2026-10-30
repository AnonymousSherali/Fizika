"""Eksport (FT-8.3): sinf jurnali (.xlsx) va topshiriq natijalari (CSV)."""
import csv
import io

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.assignments.models import Assignment, Submission
from apps.assignments.services import student_status

from .services import class_students, classroom_heatmap

STATUS_UZ = {"new": "yangi", "in_progress": "bajarilmoqda", "submitted": "topshirildi", "graded": "baholandi",
             "returned": "qaytarildi", "overdue": "muddati o'tdi"}
HEAD = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="1F6FB5")


def _style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font, cell.fill = HEAD, FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "C2"


def classroom_xlsx(classroom) -> HttpResponse:
    wb = Workbook()
    ws = wb.active
    ws.title = "Jurnal"
    assignments = list(Assignment.objects.filter(classrooms=classroom).order_by("created_at"))
    ws.append(["#", "O'quvchi"] + [f"{a.title} ({'lab' if a.type == 'lab' else 'test'})" for a in assignments]
              + ["O'rtacha (5 ballik)"])
    _style_header(ws, len(assignments) + 3)
    subs = {(s.assignment_id, s.student_id): s for s in
            Submission.objects.filter(assignment__in=assignments)}
    for i, st in enumerate(class_students(classroom), 1):
        row, fives = [i, f"{st.last_name} {st.first_name}".strip()], []
        for a in assignments:
            s = subs.get((a.id, st.id))
            if s and s.grade is not None:
                row.append(s.grade)
                if a.grading_scale == 5:
                    fives.append(s.grade)
            else:
                row.append(STATUS_UZ.get(student_status(a, s), ""))
        row.append(round(sum(fives) / len(fives), 2) if fives else None)
        ws.append(row)
    ws.column_dimensions["B"].width = 28
    for c in range(3, len(assignments) + 4):
        ws.column_dimensions[get_column_letter(c)].width = 16

    heat = classroom_heatmap(classroom)
    ws2 = wb.create_sheet("Mavzular")
    ws2.append(["#", "O'quvchi"] + [t["title"] for t in heat["topics"]] + ["Testlar soni"])
    _style_header(ws2, len(heat["topics"]) + 3)
    for i, s in enumerate(heat["students"], 1):
        ws2.append([i, s["name"]] + [s["cells"].get(t["key"]) for t in heat["topics"]] + [s["tests"]])
    ws2.column_dimensions["B"].width = 28

    buf = io.BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(),
                        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    name = f"sinf_{classroom.name}_{timezone.localdate():%Y-%m-%d}.xlsx".replace(" ", "_")
    resp["Content-Disposition"] = f'attachment; filename="{name}"'
    return resp


def assignment_csv(assignment, roster) -> HttpResponse:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";", lineterminator="\r\n")
    w.writerow(["O'quvchi", "Login", "Sinf", "Holat", "Topshirgan", "Kech", "Tavsiya", "Baho"])
    for r in roster:
        w.writerow([r["student"]["name"], r["student"]["username"], r["classroom"], STATUS_UZ.get(r["status"], r["status"]),
                    timezone.localtime(r["submitted_at"]).strftime("%d.%m.%Y %H:%M") if r["submitted_at"] else "",
                    "ha" if r["is_late"] else "", _num(r["suggested_grade"]), _num(r["grade"])])
    resp = HttpResponse("﻿" + buf.getvalue(), content_type="text/csv; charset=utf-8")
    resp["Content-Disposition"] = f'attachment; filename="topshiriq_{assignment.id}.csv"'
    return resp


def _num(v):
    return "" if v is None else (f"{v:g}".replace(".", ","))
