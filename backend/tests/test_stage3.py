"""3-bosqich: bildirishnomalar (FT-9, FT-5.8), statistika va eksport (FT-8)."""
import io
import uuid
from datetime import timedelta

import pytest
from django.core import mail
from django.utils import timezone
from openpyxl import load_workbook

from apps.assignments.models import Assignment
from apps.notifications.models import Notification
from apps.notifications.tasks import send_deadline_reminders
from apps.quizzes.models import Attempt

from .conftest import make_user
from .test_assignments import lab_payload, pendulum_journal, test_payload

pytestmark = pytest.mark.django_db


def test_new_assignment_notifies_students(api, auth, teacher, student, classroom):
    auth(api, teacher)
    api.post("/api/v1/assignments/", lab_payload(classroom), format="json")
    n = Notification.objects.get(user=student)
    assert n.type == "new_assignment" and "mayatnik_model.html?topshiriq=" in n.url
    auth(api, student)
    assert api.get("/api/v1/notifications/unread-count/").data["count"] == 1
    r = api.get("/api/v1/notifications/")
    assert r.data["count"] == 1 and r.data["results"][0]["title"].startswith("Yangi topshiriq")
    assert api.post(f"/api/v1/notifications/{n.id}/read/").data["is_read"] is True
    assert api.get("/api/v1/notifications/?unread=1").data["count"] == 0


def test_submit_grade_return_notifications(api, auth, teacher, student, classroom):
    auth(api, teacher)
    aid = api.post("/api/v1/assignments/", lab_payload(classroom), format="json").data["id"]
    auth(api, student)
    jid = pendulum_journal(api)
    sid = api.post(f"/api/v1/assignments/{aid}/submit/", {"lab_journal_id": jid, "conclusion": "x" * 30},
                   format="json").data["id"]
    assert Notification.objects.filter(user=teacher, type="submitted").count() == 1
    auth(api, teacher)
    api.post(f"/api/v1/submissions/{sid}/return/", {"comment": "Grafikni qayta chizing"}, format="json")
    api.post(f"/api/v1/submissions/{sid}/grade/", {"grade": 4}, format="json")
    types = set(Notification.objects.filter(user=student).values_list("type", flat=True))
    assert {"returned", "graded", "new_assignment"} <= types
    # O'qituvchining emaili yo'q — xat yuborilmaydi; email bo'lsa yuboriladi
    assert len(mail.outbox) == 0


def test_email_channel_and_opt_out(api, auth, teacher, classroom, school, django_capture_on_commit_callbacks):
    teacher.email = "ustoz@maktab.uz"
    teacher.save()
    s = make_user(school=school, username="emaili.bor", email="oquvchi@mail.uz")
    from apps.schools.models import ClassroomMember

    ClassroomMember.objects.create(classroom=classroom, user=s)
    auth(api, teacher)
    with django_capture_on_commit_callbacks(execute=True):
        api.post("/api/v1/assignments/", lab_payload(classroom), format="json")
    assert len(mail.outbox) == 1 and mail.outbox[0].to == ["oquvchi@mail.uz"]
    assert "mayatnik_model.html?topshiriq=" in mail.outbox[0].body
    auth(api, s)
    assert api.patch("/api/v1/me/", {"notify_email": False}, format="json").data["notify_email"] is False
    auth(api, teacher)
    with django_capture_on_commit_callbacks(execute=True):
        api.post("/api/v1/assignments/", lab_payload(classroom, title="Ikkinchi"), format="json")
    assert len(mail.outbox) == 1  # o'chirilgan — yangi xat yo'q


def test_student_join_notifies_teacher(api, classroom, teacher):
    api.post("/api/v1/auth/register/student/", {"join_code": classroom.join_code, "first_name": "Yangi",
                                                "last_name": "Oquvchi", "username": "yangi.oq", "password": "Fizika-2026!"},
             format="json")
    n = Notification.objects.get(user=teacher, type="new_student")
    assert "Oquvchi Yangi" in n.title and n.url.endswith(f"#sinf/{classroom.id}")


def test_deadline_reminders_once(api, auth, teacher, student, classroom):
    auth(api, teacher)
    due = (timezone.now() + timedelta(minutes=50)).isoformat()
    aid = api.post("/api/v1/assignments/", lab_payload(classroom, due_at=due), format="json").data["id"]
    assert send_deadline_reminders() == 2  # 24 soat va 1 soat oynalari
    assert send_deadline_reminders() == 0  # takrorlanmaydi
    # Topshirgan o'quvchiga eslatma bormaydi
    Notification.objects.filter(type="deadline").delete()
    auth(api, student)
    jid = pendulum_journal(api)
    api.post(f"/api/v1/assignments/{aid}/submit/", {"lab_journal_id": jid, "conclusion": "x" * 30}, format="json")
    assert send_deadline_reminders() == 0


def test_notifications_private(api, auth, teacher, student, classroom):
    auth(api, teacher)
    api.post("/api/v1/assignments/", lab_payload(classroom), format="json")
    n = Notification.objects.get(user=student)
    assert api.post(f"/api/v1/notifications/{n.id}/read/").status_code == 404


# ---------- Statistika ----------
def make_activity(api, auth, teacher, student, classroom):
    auth(api, teacher)
    lab = api.post("/api/v1/assignments/", lab_payload(classroom), format="json").data["id"]
    test = api.post("/api/v1/assignments/", test_payload(classroom), format="json").data["id"]
    auth(api, student)
    jid = pendulum_journal(api)
    sid = api.post(f"/api/v1/assignments/{lab}/submit/", {"lab_journal_id": jid, "conclusion": "x" * 30},
                   format="json").data["id"]
    att = Attempt.objects.get(pk=api.post("/api/v1/quiz/attempts/", {"assignment_id": test}, format="json").data["id"])
    for it in att.items.select_related("question", "problem"):
        ans = (str(it.option_order.index(it.question.correct_index)) if it.question_id else str(it.problem.answer))
        api.post(f"/api/v1/quiz/attempts/{att.id}/answer/", {"item_id": it.order, "answer": ans}, format="json")
    api.post(f"/api/v1/quiz/attempts/{att.id}/finish/")
    auth(api, teacher)
    api.post(f"/api/v1/submissions/{sid}/grade/", {"grade": 4}, format="json")
    return lab, test


def test_classroom_summary_and_heatmap(api, auth, teacher, student, classroom):
    make_activity(api, auth, teacher, student, classroom)
    r = api.get(f"/api/v1/analytics/classrooms/{classroom.id}/")
    assert r.status_code == 200
    d = r.data
    assert d["students"] == 1 and len(d["assignments"]) == 2
    lab = next(a for a in d["assignments"] if a["type"] == "lab")
    assert lab["graded"] == 1 and lab["avg_grade"] == 4
    assert d["avg_test_percent"] == 100 and d["topics"][0]["key"] == "kepler" and d["topics"][0]["percent"] == 100
    assert d["activity_7d"]["active_students"] == 1
    h = api.get(f"/api/v1/analytics/classrooms/{classroom.id}/heatmap/").data
    assert h["topics"] == [{"key": "kepler", "title": "Kepler"}]
    row = h["students"][0]
    assert row["cells"]["kepler"] == 100 and row["avg_grade"] == 4.5 and row["submitted"] == 2


def test_summary_cache_refreshes_after_grade(api, auth, teacher, student, classroom):
    lab, _ = make_activity(api, auth, teacher, student, classroom)
    first = api.get(f"/api/v1/analytics/classrooms/{classroom.id}/").data
    sub = Assignment.objects.get(pk=lab).submissions.get()
    api.post(f"/api/v1/submissions/{sub.id}/grade/", {"grade": 2}, format="json")
    second = api.get(f"/api/v1/analytics/classrooms/{classroom.id}/").data
    a1 = next(a for a in first["assignments"] if a["type"] == "lab")["avg_grade"]
    a2 = next(a for a in second["assignments"] if a["type"] == "lab")["avg_grade"]
    assert (a1, a2) == (4, 2)


def test_student_card_access(api, auth, teacher, other_teacher, student, classroom):
    make_activity(api, auth, teacher, student, classroom)
    r = api.get(f"/api/v1/analytics/students/{student.id}/")
    assert r.status_code == 200
    d = r.data
    assert d["avg_grade"] == 4.5 and len(d["assignments"]) == 2 and d["attempts"][0]["percent"] == 100
    assert d["labs"][0]["count"] == 1 and d["weak_topics"] == []
    auth(api, student)
    assert api.get(f"/api/v1/analytics/students/{student.id}/").status_code == 200
    auth(api, other_teacher)
    assert api.get(f"/api/v1/analytics/students/{student.id}/").status_code == 403
    assert api.get(f"/api/v1/analytics/classrooms/{classroom.id}/").status_code == 404


def test_classroom_xlsx_export(api, auth, teacher, student, classroom):
    make_activity(api, auth, teacher, student, classroom)
    r = api.get(f"/api/v1/analytics/classrooms/{classroom.id}/export.xlsx")
    assert r.status_code == 200 and "spreadsheetml" in r["Content-Type"]
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["Jurnal"]
    assert ws.cell(1, 2).value == "O'quvchi" and ws.cell(2, 2).value == "Aliyev Vali"
    assert ws.cell(2, ws.max_column).value == 4.5
    assert wb["Mavzular"].cell(2, 3).value == 100


def test_assignment_csv_export(api, auth, teacher, student, classroom):
    lab, _ = make_activity(api, auth, teacher, student, classroom)
    r = api.get(f"/api/v1/assignments/{lab}/export/")
    assert r.status_code == 200
    text = r.content.decode("utf-8")
    assert text.startswith("﻿O'quvchi;Login") and "Aliyev Vali;aliyev.vali" in text and ";baholandi;" in text
    auth(api, student)
    assert api.get(f"/api/v1/assignments/{lab}/export/").status_code == 403


def test_platform_stats_admin_only(api, auth, teacher, student):
    auth(api, teacher)
    assert api.get("/api/v1/analytics/platform/").status_code == 403
    admin = make_user(role="admin", username="admin" + uuid.uuid4().hex[:4])
    auth(api, admin)
    d = api.get("/api/v1/analytics/platform/").data
    assert d["users"]["teacher"] >= 1 and d["users"]["student"] >= 1
